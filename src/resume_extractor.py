"""Resume upload validation, private storage and text extraction (PDF / DOCX)."""
from __future__ import annotations

import io
import re
import uuid
import zipfile
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
ALLOWED_MIME_TYPES = {
    ".pdf": {"application/pdf", "application/x-pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
        "application/zip",
    },
}
MAX_PDF_PAGES = 30
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MIN_TEXT_CHARS = 50
MAX_TEXT_CHARS = 100_000


class ResumeError(ValueError):
    """User-facing problem with an uploaded resume."""


def validate_upload(filename: str, data: bytes, max_size_mb: int, mime_type: str | None = None) -> str:
    """Validate extension, MIME type, size and file signature. Returns the normalized extension."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeError("Only PDF and DOCX resumes are accepted.")
    if not data:
        raise ResumeError("The uploaded file is empty.")
    if len(data) > max_size_mb * 1024 * 1024:
        raise ResumeError(f"The file is larger than the {max_size_mb} MB limit.")
    if mime_type and mime_type not in ALLOWED_MIME_TYPES[ext]:
        raise ResumeError(f"File type '{mime_type}' does not match a {ext} resume.")

    if ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ResumeError("The file content is not a valid PDF.")
    else:
        if not data.startswith(b"PK\x03\x04"):
            raise ResumeError("The file content is not a valid DOCX document.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
                if "word/document.xml" not in names:
                    raise ResumeError("The file is a ZIP archive but not a Word document.")
                if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ResumeError("The DOCX file expands to an unreasonable size.")
        except zipfile.BadZipFile as exc:
            raise ResumeError("The DOCX file is corrupted.") from exc
    return ext


def save_resume(data: bytes, ext: str, storage_dir: Path) -> Path:
    """Store under a server-generated name (the client filename is never used -> no path traversal)."""
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeError("Unsupported resume type.")
    storage_dir = storage_dir.resolve()
    storage_dir.mkdir(parents=True, exist_ok=True)
    path = (storage_dir / f"{uuid.uuid4().hex}{ext}").resolve()
    if path.parent != storage_dir:
        raise ResumeError("Invalid storage path.")
    path.write_bytes(data)
    return path


def resolve_stored_path(stored: str, storage_dir: Path) -> Path:
    """Resolve a DB-stored resume path and make sure it stays inside the private storage dir."""
    storage_dir = storage_dir.resolve()
    path = Path(stored)
    if not path.is_absolute():
        path = storage_dir / path
    path = path.resolve()
    if storage_dir not in path.parents:
        raise ResumeError("Resume path is outside private storage.")
    return path


def _extract_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several error types for corrupt files
        raise ResumeError("The PDF could not be opened (corrupted file).") from exc
    with doc:
        if doc.needs_pass:
            raise ResumeError("Password-protected PDFs are not supported.")
        if doc.page_count > MAX_PDF_PAGES:
            raise ResumeError(f"Resumes longer than {MAX_PDF_PAGES} pages are not supported.")
        return "\n".join(page.get_text("text") for page in doc)


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeError("The DOCX file could not be read.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, ext: str) -> str:
    text = _extract_pdf(data) if ext == ".pdf" else _extract_docx(data)
    text = text.replace("\x00", "").strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeError("No readable text found in the resume (scanned image PDFs are not supported).")
    if len(text) > MAX_TEXT_CHARS:
        raise ResumeError("The resume text is unusually long; please upload a shorter resume.")
    return text


_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}")


def find_email(text: str) -> str | None:
    """First email address in the resume text (deterministic; no LLM needed for contact details)."""
    match = _EMAIL_RE.search(text or "")
    return match.group(0).lower() if match else None


def extract_text_from_path(path: Path) -> str:
    return extract_text(path.read_bytes(), path.suffix.lower())
